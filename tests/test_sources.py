"""Imported sources: offline parsing, provenance, indexing, lifecycle and safe URL fetching."""
from __future__ import annotations

import hashlib
import io
import json
import socket
import time
from dataclasses import replace
from types import SimpleNamespace

import fitz
import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from prism.api import create_app
from prism.assist.sections import card_segment, slice_pack
from prism.llm.gateway import LLMUnavailable
from prism.orchestrator import Pipeline
from prism.rules.engine import evaluate
from prism.schemas import ClaimType, Enquiry, EvidencePack, Segment
from prism.sources.chunking import chunks
from prism.sources.extraction import extract, safe_path, validate_file
from prism.sources.service import SourceService, import_document
from prism.sources.storage import SourceStore
from prism.sources.urls import SourceHttp, canonical_url, validate_public_url
from prism.stages.checks import CheckContext, check_ids, check_web
from prism.stages.retrieve import Retriever
from prism.store import Store
from tests.factories import make_sheet, minimal_body
from tests.fakes import FakeGateway, workflow_fake


@pytest.fixture
def service(tmp_path):
    service = SourceService(SourceStore(tmp_path / "prism.db"))
    yield service
    service.close()


def wait_job(service, job):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = service.job(job["id"])
        if job["status"] in ("ready", "failed"):
            return job
        time.sleep(0.01)
    raise AssertionError("ingestion did not complete")


def pdf_bytes(scanned=False):
    with fitz.open() as doc:
        for text in ("NIR measurement of zeolite moisture.", "Raman probe measurement on the second page."):
            page = doc.new_page()
            if not scanned:
                page.insert_text((50, 50), text)
        return doc.tobytes()


def test_pdf_and_docx_extraction_preserve_pages_and_order():
    body = pdf_bytes()
    _, mime = validate_file("probe.pdf", body)
    pages, _ = extract(body, mime)
    assert [p["page"] for p in pages] == [1, 2]
    assert "second page" in pages[1]["text"] and "NIR" in pages[0]["text"]
    doc = Document()
    doc.add_paragraph("Before table")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Raman in table"
    doc.add_paragraph("After table")
    buf = io.BytesIO()
    doc.save(buf)
    _, mime = validate_file("measurement.docx", buf.getvalue())
    pages, _ = extract(buf.getvalue(), mime)
    assert pages[0]["text"] == "Before table\nRaman in table\nAfter table"
    with pytest.raises(ValueError, match="OCR"):
        extract(pdf_bytes(scanned=True), "application/pdf")


@pytest.mark.parametrize("name", ["../probe.txt", "/probe.txt", "folder/../../probe.txt", "C:\\probe.txt",
                                  "folder\\..\\probe.txt", "folder//probe.txt", "\x00probe.txt"])
def test_path_traversal_rejected(name):
    with pytest.raises(ValueError, match="path"):
        safe_path(name)


@pytest.mark.parametrize(("name", "content"), [("bad.pdf", b"ordinary text"), ("bad.txt", b"%PDF-1.7"),
                                                ("bad.docx", b"PK not docx"), ("bad.exe", b"MZ"),
                                                ("bad.html", b"ordinary text"), ("bad.txt", b"binary\x00"),
                                                ("bad.md", b"\xff\xfe")])
def test_type_checks(name, content):
    with pytest.raises(ValueError):
        validate_file(name, content)


def test_sanitised_name_keeps_relative_path_without_using_it_for_storage(service):
    job = service.upload("default", "Probe", b"NIR document", "folder/<probe>.txt", background=False)
    source = service.detail(job["source_id"])
    assert source["relative_path"] == "folder/_probe_.txt"
    assert service.store.blob(source["id"]).parent == service.store.files
    assert service.store.blob(source["id"]).name == source["id"]


def test_chunk_spans_hashes_overlap_and_page_boundaries():
    pages = [{"page": 3, "text": "NIR moisture signal. " * 120}, {"page": 4, "text": "Raman measurement."}]
    rows = chunks("source", pages)
    for row in rows:
        page = next(p for p in pages if p["page"] == row["page"])
        assert row["text"] == page["text"][row["start"]:row["end"]]
        assert row["sha256"] == hashlib.sha256(row["text"].encode()).hexdigest()
    assert rows[1]["start"] < rows[0]["end"]
    assert rows[-1]["page"] == 4 and rows[-1]["start"] == 0


def test_duplicate_delete_reindex_and_collections(service):
    collection = service.collection("Research")
    job = service.upload(collection["id"], "Zeolite", b"Zeolite moisture spectroscopy.", "zeolite.md", background=False)
    assert job["status"] == "ready"
    source = service.detail(job["source_id"])
    assert source["chunks"] == 1 and source["pages"] == 1
    duplicate = service.upload("default", "Other name", b"Zeolite moisture spectroscopy.", "duplicate.txt")
    assert duplicate["duplicate"] and duplicate["source_id"] == source["id"]
    assert service.list()["total"] == 1
    service.collection("Renamed", collection["id"])
    assert service.list(collection["id"])["total"] == 1
    assert wait_job(service, service.reindex(source["id"]))["status"] == "ready"
    assert service.detail(source["id"])["chunks"] == 1
    service.delete_collection(collection["id"])
    assert service.detail(source["id"])["collection"] == "default"
    service.delete(source["id"])
    assert service.list()["total"] == 0 and not service.cards()
    assert not service.store.blob(source["id"]).exists()
    with service.store.db() as db:
        assert db.execute("SELECT count(*) FROM source_chunks").fetchone()[0] == 0
        assert db.execute("SELECT count(*) FROM source_embeddings").fetchone()[0] == 0
    with pytest.raises(KeyError):
        service.detail(source["id"])


def test_background_failure_visible_and_recoverable(service):
    job = service.upload("default", "Scanned", pdf_bytes(scanned=True), "scanned.pdf")
    job = wait_job(service, job)
    assert job["status"] == "failed" and "OCR" in job["error"]
    assert service.detail(job["source_id"])["status"] == "failed"
    assert service.cards() == []
    assert wait_job(service, service.reindex(job["source_id"]))["status"] == "failed"


def test_html_extraction_removes_scripts_navigation_hidden_text():
    body = b'<html><title>Probe guide</title><nav>Navigation</nav><main><h1>NIR</h1><p>Moisture.</p><p hidden>hidden</p><script>execute()</script></main></html>'
    pages, title = extract(body, "text/html")
    assert title == "Probe guide" and "Moisture." in pages[0]["text"]
    assert all(text not in pages[0]["text"] for text in ("Navigation", "execute()", "hidden"))


def resolver(host, *args, **kwargs):
    ip = "10.0.0.2" if host == "private.example" else "93.184.216.34"
    return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, 0))]


@pytest.mark.parametrize("url", ["file:///etc/passwd", "ftp://public.example/x", "http://localhost/x",
                                "http://127.0.0.1", "http://[::1]", "http://169.254.169.254/latest/meta-data/",
                                "http://10.0.0.1", "http://private.example", "http://2130706433/", "http://0x7f.1/", "https://user:pass@public.example/"])
def test_url_ssrf_guard(url):
    with pytest.raises(ValueError):
        validate_public_url(url, resolver)


def test_url_guard_rechecks_redirect_dns_and_never_contacts_private_host():
    calls = []
    def transport(request):
        calls.append(str(request.url))
        return httpx.Response(302, headers={"location": "http://private.example/metadata"})
    http = SourceHttp(httpx.MockTransport(transport), resolver)
    try:
        response = http.get("https://public.example/start", page=True)
        assert not response.ok and len(calls) == 1
    finally:
        http.close()


def test_url_import_canonical_duplicate_provenance_and_limits(service):
    def transport(request):
        if request.url.path == "/large":
            return httpx.Response(200, headers={"content-type": "text/html", "content-length": "9999999"})
        if request.url.path == "/binary":
            return httpx.Response(200, headers={"content-type": "image/png"}, content=b"binary")
        if request.url.path == "/start":
            return httpx.Response(302, headers={"location": "/guide"})
        return httpx.Response(200, headers={"content-type": "text/html"}, content=b"<html><title>Guide</title><main><p>NIR zeolite moisture.</p></main></html>")
    service.http = SourceHttp(httpx.MockTransport(transport), resolver)
    assert canonical_url("HTTPS://PUBLIC.EXAMPLE:443/start#section") == "https://public.example/start"
    job = wait_job(service, service.urls("default", ["HTTPS://PUBLIC.EXAMPLE:443/start#section"])[0])
    assert job["status"] == "ready"
    source = service.detail(job["source_id"])
    assert source["title"] == "Guide" and source["kind"] == "WEB"
    assert source["origin"]["final_url"] == "https://public.example/guide"
    assert source["origin"]["retrieved_at"]
    card = service.cards()[0]
    assert card.quote in source["preview"][0]["text"] and "Web: public.example (retrieved" in card.citation()
    assert service.urls("default", ["https://public.example/start#other"])[0]["duplicate"]
    assert wait_job(service, service.urls("default", ["https://public.example/guide"])[0])["duplicate"]
    for path in ("large", "binary"):
        assert wait_job(service, service.urls("default", [f"https://public.example/{path}"])[0])["status"] == "failed"


def test_import_document_connector_contract(service, monkeypatch):
    monkeypatch.setattr("prism.sources.service.default_service", lambda: service)
    origin = {"kind": "mcp", "server": "Research", "uri": "resource://probe", "retrieved_at": "2026-10-18T10:00:00Z"}
    result = import_document("default", "Probe resource", b"NIR moisture reference.", "text/plain", origin)
    assert result["status"] == "ready" and not result["duplicate"]
    source = service.detail(result["source_id"])
    assert source["kind"] == "MCP" and source["origin"] == origin
    card = service.cards()[0]
    assert card.source == "MCP" and card.provenance["uri"] == "resource://probe"
    assert "Imported through MCP: Probe resource" in card.citation()
    assert import_document("default", "Duplicate", b"NIR moisture reference.", "text/plain", origin)["duplicate"]
    with pytest.raises(ValueError, match="MCP origin"):
        import_document("default", "Missing provenance", b"text", "text/plain", {})


def test_imported_retrieval_refreshes_and_never_becomes_afs(kb, service):
    view = replace(kb, cards=dict(kb.cards), imported_sources=service)
    retriever = Retriever(view)
    job = service.upload("default", "Zeolite", b"Zeolite spectroscopic moisture measurement.", "zeolite.txt", background=False)
    hits = retriever.search("zeolite", k=1)
    card = view.cards[hits[0]]
    assert card.source == "UPLOAD" and card.provenance["source_id"] == job["source_id"]
    assert card.citation() == "Uploaded: zeolite p.1"
    assert card_segment(card).claim_type == ClaimType.EXTERNAL
    sheet = make_sheet(matrix="liquid", targets=("zeolite",))
    pipeline = Pipeline(workflow_fake(sheet, classifier=[{"label": "IN_SCOPE", "reason": "t"}]), view)
    enquiry = Enquiry(case_id="upload", text="zeolite", sha256="test")
    prepared = pipeline.prepare(enquiry)
    pack = retriever.pack(sheet, evaluate(sheet, view))
    assert card.id in pack.ids()
    assert all(card.id in {c.id for c in cards} for cards in slice_pack(pack, prepared.skeleton, prepared.findings, view).values())
    body = minimal_body()
    body.decision_basis = [Segment(text=card.quote, claim_type=ClaimType.EXTERNAL, evidence_ids=[card.id])]
    cx = CheckContext(body, sheet, prepared.findings, EvidencePack(cards=[card]), view, ["zeolite"])
    assert not [i for i in check_ids(cx) if "external claim" in i.message]
    assert not check_web(cx)
    body.decision_basis[0].claim_type = ClaimType.AFS
    assert any("AFS claim" in i.message for i in check_ids(cx))
    assert any("EXTERNAL" in i.message for i in check_web(cx))
    service.delete(job["source_id"])
    retriever.search("zeolite")
    assert card.id not in view.cards


def test_incremental_embeddings_do_not_reembed_existing_chunks(kb, service, tmp_path):
    class RecordingGateway(FakeGateway):
        def __init__(self):
            super().__init__(vectors={})
            self.embedded = []
        def embed(self, texts):
            self.embedded.extend(texts)
            return super().embed(texts)
    gateway = RecordingGateway()
    service.gateway = gateway
    first = service.upload("default", "First", b"First zeolite moisture source.", "a.txt", background=False)
    count = len(gateway.embedded)
    view = replace(kb, cards=dict(kb.cards), imported_sources=service)
    Retriever(view, gateway, tmp_path / "index").search("zeolite")
    first_text = gateway.embedded[0]
    assert gateway.embedded.count(first_text) == 1
    before = len(gateway.embedded)
    service.upload("default", "Second", b"Second source of zeolite spectra.", "b.txt", background=False)
    Retriever(view, gateway, tmp_path / "index").search("zeolite")
    assert count == 1 and len(gateway.embedded) - before == 2  # new chunk and query only
    assert gateway.embedded.count(first_text) == 1
    gateway.profile = SimpleNamespace(name="other", roles={"embedder": SimpleNamespace(models=["other-model"])})
    Retriever(view, gateway, tmp_path / "index").search("zeolite")
    assert gateway.embedded.count(first_text) == 2
    service.delete(first["source_id"])
    with service.store.db() as db:
        assert db.execute("SELECT count(*) FROM source_embeddings e JOIN source_chunks c ON c.id=e.chunk_id "
                          "WHERE c.source_id=?", (first["source_id"],)).fetchone()[0] == 0


@pytest.fixture
def client(kb, tmp_path):
    def factory(deep, emit, **kwargs):
        return Pipeline(workflow_fake(), kb, emit=emit)
    app = create_app(factory, Store(tmp_path / "api.db"), kb)
    with TestClient(app) as client:
        yield client


def test_sources_api_routes_and_per_file_validation(client):
    collection = client.post("/api/sources/collections", json={"name": "Documents"}).json()
    assert client.patch(f"/api/sources/collections/{collection['id']}", json={"name": "Lab"}).json()["name"] == "Lab"
    response = client.post("/api/sources/upload", data={"collection": collection["id"], "relative_paths": json.dumps(["folder/probe.txt", "../escape.txt", "fake.pdf"])},
                           files=[("files", ("probe.txt", b"NIR probe spectrometer.")), ("files", ("escape.txt", b"x")), ("files", ("fake.pdf", b"text"))])
    assert response.status_code == 202
    jobs = response.json()["jobs"]
    assert [j["status"] for j in jobs[1:]] == ["failed", "failed"]
    service = client.app.state.sources
    completed = wait_job(service, jobs[0])
    assert completed["status"] == "ready"
    sid = completed["source_id"]
    assert client.get(f"/api/sources/jobs/{jobs[0]['id']}").json()["progress"] == 100
    listed = client.get(f"/api/sources/sources?collection={collection['id']}&limit=1").json()
    assert listed["total"] == 1 and listed["items"][0]["relative_path"] == "folder/probe.txt"
    assert client.get("/api/sources/sources?offset=1").json()["items"] == []
    assert client.get(f"/api/sources/sources/{sid}").json()["preview"][0]["text"] == "NIR probe spectrometer."
    assert client.post(f"/api/sources/sources/{sid}/reindex").status_code == 202
    service._executor.shutdown(wait=True)
    # A fresh service models a restart and the sources still exist.
    service._executor = __import__('concurrent.futures').futures.ThreadPoolExecutor(max_workers=2)
    assert client.post("/api/sources/reindex").status_code == 202
    assert client.post("/api/sources/urls", json={"urls": ["file:///etc/passwd"]}).json()["jobs"][0]["status"] == "failed"
    assert client.get("/api/sources/sources/missing").status_code == 404
    assert client.get("/api/sources/sources?limit=101").status_code == 422
    assert client.delete("/api/sources/collections/default").status_code == 422
    assert client.delete(f"/api/sources/collections/{collection['id']}").status_code == 204
    assert client.delete(f"/api/sources/sources/{sid}").status_code == 204
    assert client.get(f"/api/sources/sources/{sid}").status_code == 404


def test_api_upload_batch_contract(client):
    assert client.post("/api/sources/upload", data={"relative_paths": '["one.txt","two.txt"]'}, files={"files": ("one.txt", b"one")}).status_code == 422
    assert client.post("/api/sources/upload", data={"collection": "missing"}, files={"files": ("one.txt", b"one")}).status_code == 404
    too_many = [("files", (f"{i}.txt", b"one")) for i in range(101)]
    assert client.post("/api/sources/upload", files=too_many).status_code == 413


def test_import_document_keeps_the_connector_title_and_registered_service(service):
    from prism.sources.service import set_default_service
    set_default_service(service)
    origin = {"kind": "mcp", "server": "Research", "uri": "resource://page", "retrieved_at": "2026-10-18T10:00:00Z"}
    html = b"<html><title>Page title</title><main><p>Raman probe notes.</p></main></html>"
    result = import_document("default", "Connector name", html, "text/html; charset=utf-8", origin)
    assert service.detail(result["source_id"])["title"] == "Connector name"
    with pytest.raises(ValueError, match="MIME"):
        import_document("default", "x", b"x", "application/zip", origin)


def test_file_size_limit_is_enforced():
    with pytest.raises(ValueError, match="50 MB"):
        validate_file("big.txt", b"a" * (50 * 1024 * 1024 + 1))


@pytest.mark.parametrize("title", ["", "a", "NIR", "Title Lorem Ipsum", "LOREM IPSUM", "Untitled",
                                 "Title", "Microsoft Word - Probe guide", "PowerPoint Presentation", "Slide 1"])
def test_placeholder_document_titles_fall_back_to_filename(service, title):
    with fitz.open(stream=pdf_bytes(), filetype="pdf") as pdf:
        pdf.set_metadata({"title": title})
        body = pdf.tobytes()
    job = service.upload("default", "Probe guide.pdf", body, "lab/Probe guide.pdf", background=False)
    assert service.detail(job["source_id"])["title"] == "Probe guide"


def test_meaningful_document_title_is_retained(service):
    with fitz.open(stream=pdf_bytes(), filetype="pdf") as pdf:
        pdf.set_metadata({"title": "Moisture spectroscopy guide"})
        body = pdf.tobytes()
    job = service.upload("default", "guide.pdf", body, "guide.pdf", background=False)
    assert service.detail(job["source_id"])["title"] == "Moisture spectroscopy guide"


@pytest.mark.parametrize("failure", [RuntimeError("provider outage"), TimeoutError("timeout"),
                                    LLMUnavailable("quota 429")])
def test_embedding_failure_keeps_upload_ready_and_retrievable_then_recovers(service, kb, tmp_path, failure):
    class UnavailableGateway(FakeGateway):
        def embed(self, texts):
            raise failure
    gateway = UnavailableGateway()
    service.gateway = gateway
    job = wait_job(service, service.upload("default", "Zeolite.txt", b"Zeolite moisture spectroscopic reference.", "Zeolite.txt"))
    assert job["status"] == "ready" and job["progress"] == 100
    assert job["error"] == "Semantic index pending: the embedding service was unavailable"
    source = service.detail(job["source_id"])
    assert source["status"] == "ready" and source["chunks"] and source["error"] == job["error"]
    view = replace(kb, cards=dict(kb.cards), imported_sources=service)
    retriever = Retriever(view, gateway, tmp_path / "index")
    assert service.cards()[0].id in retriever.search("zeolite", k=1)
    # The same retriever retries missing embeddings on the next search.
    gateway.embed = FakeGateway(vectors={}).embed
    retriever.search("zeolite")
    with service.store.db() as db:
        assert db.execute("SELECT count(*) FROM source_embeddings").fetchone()[0] == source["chunks"]
    assert service.detail(source["id"])["error"] == ""
    assert service.job(job["id"])["error"] == ""


def test_embedding_pending_reindex_url_and_mcp(service, monkeypatch):
    class UnavailableGateway(FakeGateway):
        def embed(self, texts):
            raise LLMUnavailable("missing key with private provider details")
    service.gateway = UnavailableGateway()
    service.http = SourceHttp(httpx.MockTransport(lambda _: httpx.Response(
        200, headers={"content-type": "text/html"}, content=b"<html><title>URL guide</title><p>Moisture reference.</p></html>")), resolver)
    web = wait_job(service, service.urls("default", ["https://public.example/guide"])[0])
    monkeypatch.setattr("prism.sources.service.default_service", lambda: service)
    mcp = import_document("default", "MCP guide", b"Connector Raman reference.", "text/plain",
                          {"kind": "mcp", "server": "Research", "uri": "resource://page", "retrieved_at": "2026-10-18"})
    for job in (web, mcp):
        assert job["status"] == "ready" and job["error"].startswith("Semantic index pending:")
        assert "private provider" not in job["error"]
    service.gateway = FakeGateway(vectors={})
    for job in (web, mcp):
        assert wait_job(service, service.reindex(job["source_id"]))["error"] == ""
        assert service.detail(job["source_id"])["error"] == ""


def test_api_source_filters_are_combined_validated_and_paginated(client):
    service = client.app.state.sources
    collection = service.collection("Research")
    ready = service.upload(collection["id"], "Ready", b"NIR spectroscopy.", "ready.txt", background=False)
    service.upload("default", "Scan", pdf_bytes(scanned=True), "scan.pdf", background=False)
    service.upload("default", "MCP", b"Connector spectra.", "resource.txt",
                   {"kind": "mcp", "server": "Lab", "uri": "resource://r"}, background=False)
    assert client.get("/api/sources/sources?status=failed").json()["total"] == 1
    assert client.get("/api/sources/sources?status=ready&kind=MCP").json()["total"] == 1
    assert client.get("/api/sources/sources?status=failed&kind=MCP").json()["total"] == 0
    listed = client.get(f"/api/sources/sources?status=ready&kind=UPLOAD&collection={collection['id']}&limit=1").json()
    assert listed["total"] == 1 and listed["items"][0]["id"] == ready["source_id"]
    assert client.get("/api/sources/sources?status=ready&offset=2").json()["items"] == []
    for query in ("status=bad", "kind=bad", "kind=MCP_LIVE"):
        assert client.get("/api/sources/sources?" + query).status_code == 422


def test_api_folder_skips_unsupported_and_system_files_and_ignores_missing_mime(client):
    names = ["lab/nested/notes.md", "lab/notes.txt", "lab/page.html", "lab/.DS_Store", "lab/Thumbs.db", "lab/data.csv", "lab/.hidden/notes.txt"]
    bodies = [b"Moisture notes.", b"NIR reference.", b"<html><p>Raman.</p></html>", b"binary", b"binary", b"a,b", b"hidden"]
    response = client.post("/api/sources/upload", data={"relative_paths": json.dumps(names)},
                           files=[("files", (name.split("/")[-1], body, "")) for name, body in zip(names, bodies)])
    assert response.status_code == 202
    data = response.json()
    assert len(data["jobs"]) == 3 and len(data["skipped"]) == 4
    assert [item["path"] for item in data["skipped"]] == names[3:]
    assert all(item["reason"] for item in data["skipped"])
    for job in data["jobs"]:
        assert wait_job(client.app.state.sources, job)["status"] == "ready"
    assert client.get("/api/sources/sources").json()["total"] == 3
